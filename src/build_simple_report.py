"""用原课程模板制作单份小组报告，内容与simple_predict.py一致。"""
import json
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from build_reports import ROOT, TEAM, clear, para, table, font, preserve_save

reference = ROOT / '实验报告模板.docx'
doc = Document(reference)
summary = json.loads((ROOT / 'simple_results/summary.json').read_text(encoding='utf-8'))
t = doc.tables[0]
t.cell(0, 1).text = '地铁卡介质客流预测'
t.cell(1, 1).text = '；'.join(f'{n} {s}' for n, s, _, _ in TEAM)
t.cell(2, 1).text = '智慧交通23-1；交通23'
for row in t.rows:
    for c in row.cells:
        for p in c.paragraphs:
            for r in p.runs: font(r)
font(doc.paragraphs[2].add_run('报告形式：小组共同报告'), 10.5)
cells = [t.cell(0, 0) for t in doc.tables[1:5]]
for t in doc.tables[1:5]:
    for row in t.rows:
        row.height = None
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for e in row._tr.xpath('./w:trPr/w:cantSplit'): e.getparent().remove(e)
for c in cells: clear(c)
c = cells[0]
para(c, '1.1 实验目标', True, 12)
para(c, '利用固定的地铁日票务数据，预测智能卡与NCMC卡介质合计客流。实验以一个可直接运行的Python文件完成CSV读取、数据检查、历史预测、误差计算及结果保存，理解交通数据处理和时间序列预测的基本流程。预测单位为人次/日，不代表独立乘客人数，也不包含所有票种。')
para(c, '1.2 实验意义与分工', True, 12)
para(c, '日客流预测可为票务服务准备和需求趋势分析提供参考。当前数据没有站点、小时或运能字段，因此不直接用于车站拥挤判断及列车排班。简化实现采用同星期历史均值，算法透明、无需第三方依赖，便于课堂运行和解释。')
table(c, ['姓名', '学号', '分工安排'], [(n,s,role) for n,s,_,role in TEAM], [2.5,4.2,8.5])
c = cells[1]
para(c, '2.1 数据来源与准备', True, 12)
para(c, '数据集为Kaggle的Bengaluru NammaMetro Ridership Dataset，作者Mahesh Shantaram，使用固定第45版。原作者从BMRCL公开客流页面采集。CSV包含460条日记录、14列，日期从2024年10月26日至2026年8月16日。660个自然日中缺少200天，缺失日期不补成零客流。原始历史记录未逐行与运营方归档核对。')
para(c, '使用Record Date解析日期，将Total Smart Cards与Total NCMC相加作为目标。程序检查非负整数计数、非正合计及重复日期；发现异常则报错。此前审计发现QR汇总与子项有218行不一致，因此本简化实验不使用QR列，也不将所有票种列直接相加。')
para(c, '2.2 预测方法', True, 12)
para(c, '采用过去四周同星期的可用客流均值。预测日为d时，查询d−7、d−14、d−21、d−28四个日历日期，取存在记录的值计算均值。若均缺失，退回d−29至d−2之间的历史均值；仍无记录则用最近可用值。历史数据距离预测日超过7天时停止预测。')
para(c, '假设在预测日前一天发出结果，并考虑约一天的数据发布延迟，程序最多使用d−2及更早的观测。发布延迟是实验假设，原数据未提供每行真实发布时间。算法直接使用历史均值，无需训练参数，也不依赖pandas或scikit-learn。')
para(c, '2.3 运行步骤', True, 12)
para(c, '安装Python 3.8或更高版本。将simple_predict.py与data文件夹放在同一目录，保留data/raw/NammaMetro_Ridership_Dataset.csv。进入该目录运行：python simple_predict.py。程序打印结果，并生成simple_results/predictions.csv和simple_results/summary.json。整个过程离线运行，不需要联网下载数据。')
para(c, '程序按日期排序，将最后75个有记录日作为历史检验期；逐日预测只使用当时可获得的历史值。滚动检验允许使用前面测试日已经发布的观测，并不是一次性预测未来75天。短数据时测试数量最多为样本数的20%，本固定数据实际为75天。')
c = cells[2]
para(c, '3.1 历史检验结果', True, 12)
para(c, '测试日期为2026年5月1日至8月16日，共75个有真实记录的日期，缺失日不评分。以下结果来自实际运行simple_predict.py，不是手工填入的模拟数据。')
table(c, ['指标', '结果'], [('原始观测数','460条'),('测试样本数','75日'),('MAE',f"{summary['MAE']:,.2f}人次/日"),('MAPE',f"{summary['MAPE_percent']:.2f}%"),('预测方法','过去四周同星期均值')], [6,9.2])
para(c, 'MAE是预测值与真实值绝对差的平均值，表示平均每日偏差约2.90万人次。MAPE是每日绝对误差除以当日真实客流后取平均，约7.96%；不能把它直接称为92.04%的预测准确率。')
para(c, '3.2 数据末端预测', True, 12)
para(c, f"数据最后日期为2026年8月16日。按发布延迟假设，对8月18日生成预测：{summary['prediction']:,.2f}人次，约60.93万人次。该值是固定历史快照下的演示结果，未提供该日真实值，因此不能评价这一条预测的误差，也不能称为当前日期的实时预测。简化代码只输出点预测，不输出分类、聚类或预测区间。")
para(c, '3.3 失败案例', True, 12)
para(c, '2026年5月1日真实客流为198232人次，预测为515221.67人次，绝对误差为316989.67人次，相对误差约159.91%。这表明同星期均值不能充分应对特殊日期的需求变化。本实验保留该日参与评分，没有通过删除异常日期美化指标；仅凭日期也不能确定偏差的具体原因。')
c = cells[3]
para(c, '4.1 实验总结', True, 12)
para(c, '本实验用单文件程序实现了可运行的静态交通数据预测流程。日期排序、按日历查找历史、缺失日不填零及限制未来信息，是保证结果可解释的关键。简单的同星期均值能够利用周周期，并为后续复杂模型提供可复核的比较基线。')
para(c, '4.2 局限与改进', True, 12)
para(c, '数据仅覆盖一个城市的部分票种，缺失日较多，缺失机制可能影响评价。模型未考虑节假日、天气或运营变化，无法保证特殊日期仍准确。后续可补齐可信事件数据，并在新的未见时间窗口比较改进模型。本次简化交付保留已有完整项目作为历史资料，报告结果以simple_predict.py为准。')
para(c, '4.3 可复现材料与资料来源', True, 12)
para(c, '代码：simple_predict.py；数据：data/raw/NammaMetro_Ridership_Dataset.csv；逐日结果：simple_results/predictions.csv；汇总：simple_results/summary.json。分工按小组提供的信息列示，真实操作记录、个人心得及签名由成员核实。')
para(c, '[1] 北方工业大学课程资料：https://gitee.com/itsncut/datascience/tree/master/practiceclass')
para(c, '[2] Kaggle数据集：https://www.kaggle.com/datasets/maheshshantaram/bengaluru-nammametro-daily-ridership （第45版，CC0）。')
para(c, '[3] BMRCL公开客流：https://www.bmrc.co.in/ridership/ 。原作者采集程序：https://github.com/thecont1/namma-metro-ridership-tracker 。')
output = ROOT/'reports/地铁卡介质客流预测_小组报告_简化版.docx'
print(preserve_save(doc, reference, output))
